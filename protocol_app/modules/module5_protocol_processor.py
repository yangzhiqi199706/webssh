# -*- coding: utf-8 -*-
"""模块5：协议文件批量处理器"""
import os
import pandas as pd
import re
import shutil
from datetime import datetime
import json
from docx import Document

class ProtocolProcessor:
    """协议文件处理器"""
    
    def __init__(self, upload_folder, output_folder):
        self.upload_folder = upload_folder
        self.output_folder = output_folder
        self.processing_log = []
        
        # 确保输出文件夹存在
        os.makedirs(self.output_folder, exist_ok=True)
    
    def allowed_file(self, filename):
        """验证文件格式"""
        allowed_extensions = {'xls', 'xlsx', 'docx'}
        return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_extensions
    
    def normalize_filename(self, original_filename):
        """规范化文件名"""
        # 移除版本号和括号
        name = re.sub(r'\([^)]*\)', '', original_filename)
        name = re.sub(r'\[[^\]]*\]', '', name)
        name = re.sub(r'\d{8}\(?\d*\)?', '', name)  # 移除日期版本号
        name = re.sub(r'\(?\d+\)?', '', name)  # 移除数字版本号
        name = re.sub(r'_+', '_', name)  # 合并多个下划线
        name = name.strip('_').strip()
        
        # 确保扩展名为.xlsx
        if name.lower().endswith('.xls'):
            name = name[:-4] + '.xlsx'
        elif name.lower().endswith('.docx'):
            name = name[:-5] + '.xlsx'
        elif not name.lower().endswith('.xlsx'):
            name += '.xlsx'
        
        return name
    
    def clean_nan_values(self, df):
        """清洗NaN值"""
        cleaned_count = 0
        for col in df.columns:
            nan_count = df[col].isna().sum()
            if nan_count > 0:
                df[col] = df[col].fillna("")
                cleaned_count += nan_count
        
        return df, cleaned_count
    
    def process_protocol_file(self, file_path, original_filename):
        """处理单个协议文件"""
        try:
            self.log(f"开始处理: {original_filename}")
            
            # 读取文件
            if original_filename.lower().endswith('.docx'):
                # 处理Word文档
                df = self.read_word_to_excel(file_path)
                if df is None:
                    self.log(f"Word文档处理失败: {original_filename}")
                    return None, "Word文档解析失败"
            else:
                # 处理Excel文件
                df = pd.read_excel(file_path, header=None)
            
            # 清洗NaN值
            df, cleaned_count = self.clean_nan_values(df)
            
            # 规范化文件名
            normalized_name = self.normalize_filename(original_filename)
            output_path = os.path.join(self.output_folder, normalized_name)
            
            # 保存处理后的文件
            df.to_excel(output_path, index=False, header=False, engine='openpyxl')
            
            result = {
                "original_name": original_filename,
                "normalized_name": normalized_name,
                "rows": int(len(df)),
                "columns": int(len(df.columns)),
                "cleaned_nan": int(cleaned_count),
                "status": "success"
            }
            
            self.log(f"处理完成: {normalized_name} ({len(df)}行, {len(df.columns)}列)")
            return result, None
            
        except Exception as e:
            error_msg = f"处理失败: {str(e)}"
            self.log(error_msg)
            return None, error_msg
    
    def read_word_to_excel(self, file_path):
        """将Word文档转换为DataFrame"""
        try:
            doc = Document(file_path)
            
            # 提取所有表格数据
            tables_data = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        row_data.append(cell_text if cell_text else "")
                    table_data.append(row_data)
                tables_data.extend(table_data)
            
            # 如果没有表格，尝试提取段落文本
            if not tables_data:
                paragraphs_data = []
                for paragraph in doc.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        paragraphs_data.append([text])
                
                if paragraphs_data:
                    df = pd.DataFrame(paragraphs_data)
                    self.log(f"Word文档段落提取成功: {len(df)}行")
                else:
                    self.log("Word文档中未找到表格或文本内容")
                    return None
            else:
                df = pd.DataFrame(tables_data)
                self.log(f"Word文档表格提取成功: {len(df)}行, {len(df.columns)}列")
            
            # 确保DataFrame不为空
            if df.empty:
                self.log("Word文档提取结果为空")
                return None
            
            return df
            
        except Exception as e:
            error_detail = str(e)
            self.log(f"Word文档读取失败: {error_detail}")
            
            # 尝试提供更详细的错误信息
            import traceback
            self.log(f"详细错误: {traceback.format_exc()}")
            return None
    
    def batch_process(self, file_list):
        """批量处理文件"""
        results = []
        success_count = 0
        fail_count = 0
        
        self.log(f"开始批量处理，共{len(file_list)}个文件")
        
        for file_info in file_list:
            file_path = file_info.get('path')
            original_filename = file_info.get('name')
            
            if not file_path or not original_filename:
                continue
            
            result, error = self.process_protocol_file(file_path, original_filename)
            
            if result:
                results.append(result)
                success_count += 1
            else:
                results.append({
                    "original_name": original_filename,
                    "status": "failed",
                    "error": error
                })
                fail_count += 1
        
        summary = {
            "total": len(file_list),
            "success": success_count,
            "failed": fail_count,
            "results": results
        }
        
        self.log(f"批量处理完成: 成功{success_count}个, 失败{fail_count}个")
        return summary
    
    def log(self, message):
        """记录日志"""
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        log_entry = f"[{timestamp}] {message}"
        self.processing_log.append(log_entry)
        print(log_entry)
    
    def get_processing_log(self):
        """获取处理日志"""
        return self.processing_log
    
    def get_available_files(self, folder_path):
        """获取文件夹中的可用文件"""
        files = []
        if os.path.exists(folder_path):
            for filename in os.listdir(folder_path):
                file_path = os.path.join(folder_path, filename)
                if os.path.isfile(file_path) and self.allowed_file(filename):
                    files.append({
                        "name": filename,
                        "path": file_path,
                        "size": int(os.path.getsize(file_path))
                    })
        return files
    
    def generate_report(self, summary):
        """生成处理报告"""
        report_lines = [
            "协议文件批量处理报告",
            "=" * 50,
            f"处理时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"总文件数: {summary['total']}",
            f"成功处理: {summary['success']}",
            f"处理失败: {summary['failed']}",
            "=" * 50,
            ""
        ]
        
        for result in summary['results']:
            if result['status'] == 'success':
                report_lines.append(f"✓ {result['original_name']}")
                report_lines.append(f"  → {result['normalized_name']}")
                report_lines.append(f"  → {result['rows']}行, {result['columns']}列")
                report_lines.append(f"  → 清洗了{result['cleaned_nan']}个NaN值")
            else:
                report_lines.append(f"✗ {result['original_name']}")
                report_lines.append(f"  → 错误: {result['error']}")
            report_lines.append("")
        
        report_lines.extend([
            "=" * 50,
            "处理日志:",
            ""
        ])
        
        report_lines.extend(self.processing_log)
        
        return "\n".join(report_lines)